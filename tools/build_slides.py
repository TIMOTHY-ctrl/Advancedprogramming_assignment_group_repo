"""Build the 15-slide coursework deck with editable PowerPoint text and a printable browser version."""
from html import escape
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SLIDES = [
    ("Bus seat booking", "DDD · TDD · Clean Architecture", [
        "ATUHAIRE MARY SEANICE · 2300706843 — Domain Lead; Slides 1–3",
        "kisakye rita · 2400700571 — Domain Model Lead; Slides 4–6",
        "Nassaka Catherine · 2400709816 — Architecture Lead; Slides 7–9",
        "Mwizerwa Timothy · 2400724330 — Testing Lead; Slides 11–12",
        "Mwizerwa Timothy also covers Integration Lead; Slides 10 and 13–15",
        "Makerere University · Advanced Programming group coursework",
    ]),
    ("A seat must be sold once", "Problem · scope · language", [
        "Problem: reserve bus seats without duplicate bookings, and stop advertising a trip when its final seat is sold.",
        "Inside scope: schedule a trip, book a seat, retrieve a trip, list active advertisements, and withdraw an advertisement through an event.",
        "Outside scope: payments, cancellations, journey durations, overlapping journeys, web UI, database, and distributed messaging.",
        "Trip = one scheduled departure; Bus = vehicle identified by bus number; Booking = identified seat reservation.",
        "SeatNumber = positive integer value; TripAdvertisement = public availability record identified by trip number.",
        "Two connected use cases: BookSeat → WithdrawAdvertisement. Console inputs use local dates and times.",
    ]),
    ("Exactly six business rules", "Rule · responsible component · violation outcome", [
        "BR1 — Value: a seat number must be a positive integer, excluding bool. SeatNumber rejects invalid values with ValueError.",
        "BR2 — Identity/state: booking an available seat adds an identified Booking to the same Trip and reduces availability. Trip.book_seat rejects empty passenger names or a full trip.",
        "BR3 — Invariant: each booked seat is unique and within 1..bus capacity. Trip.book_seat rejects duplicates/out-of-range seats; its collection is read-only to callers.",
        "BR4 — Cross-concept: the same bus cannot serve two trips at the same departure instant. BusSchedulingPolicy checks existing trips and rejects the proposed assignment before storage.",
        "BR5 — Follow-up: the final booking raises TripFullyBooked to request withdrawal of that trip's advertisement. TripAdvertisement permits active → withdrawn only; repeated withdrawal returns a rejection with no state change.",
        "BR6 — Lookup: booking requires an existing Trip; scheduling requires an unused trip number. Application services use ITripRepository and reject missing/duplicate identities before mutation.",
    ]),
    ("Values and identities", "BR1 and BR2", [
        "SeatNumber has no separate identity: SeatNumber(1) equals another SeatNumber(1). It is immutable and validates its value at construction.",
        "Trip is identified by trip_number, which remains stable as bookings change.",
        "Booking is a child entity identified by a generated booking_id; passenger name is data, not identity.",
        "Bus is identified operationally by bus_number; the scheduling policy compares that identity across trips.",
        "Frozen dataclasses prevent accidental field reassignment. Identity in this model is the explicit ID, not dataclass structural equality.",
    ]),
    ("Two roots protect two boundaries", "BR3 · aggregate invariants", [
        "Aggregate A root: Trip. Contents: Bus assignment and Booking children, each containing a SeatNumber.",
        "A invariant: no duplicate seat, every booked seat lies within capacity, and available_seats = capacity − booking count ≥ 0.",
        "Trip creates bookings; callers receive a read-only view of the booking collection. Booking children have no independent repository.",
        "Aggregate B root: TripAdvertisement. Identity: trip_number. No child entities are needed for this small aggregate.",
        "B invariant: an active advertisement has positive availability; a withdrawn advertisement has zero availability. Withdrawal is one-way.",
        "B returns new immutable state from update_availability/withdraw. Only its own methods determine valid transitions.",
    ]),
    ("A policy across trip assignments", "BR4 · Domain Service · Factory decision", [
        "BusSchedulingPolicy receives a proposed Bus, departure time, and existing Trip aggregates.",
        "It rejects equal departure times for an existing assignment with the same bus_number. Another bus at the same instant remains valid.",
        "This belongs in a Domain Service: the decision spans the proposed assignment and other Trip roots; no individual Trip owns the schedule.",
        "CreateTripService loads existing trips and invokes the policy before adding the new Trip or advertisement.",
        "No Factory: construction is small and constructors already validate inputs. A Factory would add indirection without a creation workflow.",
        "Deliberate limit: no journey-duration model, so this rule does not detect overlapping journeys at different departure times.",
    ]),
    ("Repositories store aggregate roots", "Contracts in Application · adapters in Infrastructure", [
        "ITripRepository → InMemoryTripRepository: stores Trip roots keyed by trip_number, including their Booking children.",
        "Operations: add, find_by_number, list_all, save. Duplicate identities and saving unknown trips are rejected.",
        "ITripAdvertisementRepository → InMemoryTripAdvertisementRepository: stores TripAdvertisement roots keyed by trip_number.",
        "Operations: put, find_by_number, list_available. Withdrawn roots are retained; available listings include active roots only.",
        "One abstraction per aggregate root used in the workflow. No separate Booking repository breaks the Trip boundary.",
        "Both implementations are process-local dictionaries. Restarting the console resets the data.",
    ]),
    ("Dependencies point toward policy", "Static import arrows: caller → dependency", [
        "Interface → Application contracts / DTOs; Interface → Domain Bus",
        "Application → Domain entities / events / policy",
        "Infrastructure → Application repository contracts; Infrastructure → Domain roots",
        "Domain → Python standard library and other Domain modules only",
        "main.py is the composition root → Interface + Application + Infrastructure + Domain",
        "No Application → Infrastructure import. Repository implementations enter through constructor injection.",
        "No Layer Supertype: entities have no shared behaviour needing a common base. Repository ABCs are contracts, not a domain Layer Supertype.",
    ]),
    ("BookSeatService coordinates", "Input/output DTOs · dependency injection", [
        "Input: BookSeatInputDTO(trip_number, passenger_name, seat_number). The console parses user input; the domain validates seat values.",
        "Service flow: repository lookup → SeatNumber → Trip.book_seat → save Trip → refresh/handle advertisement → output DTO.",
        "Output: booking_id, trip_number, passenger_name, seat_number, available_seats, advertisement_outcome.",
        "Business decisions stay in SeatNumber, Trip, TripAdvertisement, and BusSchedulingPolicy.",
        "main.build_console_app constructs the two in-memory repositories and injects them into services/handler. Tests inject fresh repositories.",
        "A rejected follow-up is exposed explicitly. The booking remains committed; this example does not promise cross-aggregate transaction rollback.",
    ]),
    ("The final seat triggers withdrawal", "BR5 · in-process aggregate communication", [
        "Request → BookSeatService → Trip (A) → TripFullyBooked → TripFullyBookedHandler → TripAdvertisement (B)",
        "Before the event: a valid booking reduces available_seats to zero. Earlier bookings produce no event.",
        "The application saves Trip and passes the event to the handler in the same process; no broker is needed.",
        "The handler retrieves B and calls withdraw(). B checks that it is active, then returns inactive state with zero availability.",
        "The handler stores the returned root; listings hide it. It returns 'withdrawn' to the booking DTO.",
        "Repeated delivery: B rejects withdrawal; the handler returns 'rejected: advertisement already withdrawn'. State is unchanged. Missing B also returns a rejection.",
    ]),
    ("Eight tests, explicit coverage", "test/test_architecture.py · expected results", [
        "T1 / BR1: seat 1 accepted; zero, negative, bool, float, and string rejected.",
        "T2 / BR2: booking ID stored on the identified Trip; availability falls from 2 to 1.",
        "T3 / BR3: seat equal to capacity accepted; duplicate and capacity + 1 rejected without extra bookings.",
        "T4 / BR4: same bus/time rejected before storage; different bus/time accepted.",
        "T5 / BR5: no event before final seat; final booking returns TripFullyBooked with the trip identity.",
        "T6 / BR6: missing Trip and duplicate trip number rejected; existing state unchanged.",
        "T7: full BookSeat use case returns 'withdrawn'; B is inactive/zero and absent from available listings.",
        "T8: repeated follow-up rejected; both final aggregate states and listing stay unchanged. Actual suite result: 8 tests, OK.",
    ]),
    ("A real fail → implement → pass cycle", "T4 · captured during this revision", [
        "RED: wrote T4 before the scheduling policy existed and ran it against the original scheduling behaviour (imports repaired).",
        "Actual failure: AssertionError: ValueError not raised. Evidence: evidence/tdd-red.txt (one test, one failure).",
        "IMPLEMENT: added BusSchedulingPolicy.ensure_available and invoked it before CreateTripService stores a Trip.",
        "GREEN: reran the same T4 unchanged; it passed. Evidence: evidence/tdd-green.txt.",
        "REFACTOR / REGRESSION: final complete suite contains exactly eight tests and passes. Evidence: evidence/test-results.txt.",
        "Inspect evidence/tdd-change.diff for the relevant change. This documents the new BR4 cycle, not the original project's development history.",
    ]),
    ("Walk through the main use case", "TR001 · two seats · final booking", [
        "Setup: schedule TR001 on BUS001 with capacity 2. The scheduling policy accepts; both Trip and its active advertisement are stored.",
        "Alex books seat 1. Trip availability becomes 1; no full-booking event. The advertisement updates to 1.",
        "Sam submits BookSeatInputDTO('TR001', 'Sam', 2). Lookup succeeds and SeatNumber(2) validates.",
        "Trip verifies range, capacity, uniqueness, and passenger name. It creates Booking, leaves zero seats, and raises TripFullyBooked('TR001').",
        "Application saves Trip. Handler asks B to withdraw; B accepts and returns its inactive state, which is stored.",
        "Output reports the booking ID, zero remaining seats, and 'withdrawn'. Available-trip listing excludes TR001. T7 verifies this complete path.",
    ]),
    ("Change a design when evidence demands it", "Rejected choice · improved responsibility", [
        "Rejected: the event handler directly deletes the advertisement through a repository. B then has no transition rule and cannot reject a follow-up.",
        "Chosen: retain B and let TripAdvertisement.withdraw guard the state change. The handler orchestrates and persists its result.",
        "Benefit: T8 can demonstrate a real domain rejection and unchanged final state. Retention also makes the outcome inspectable.",
        "Tradeoff: inactive records remain in memory; list_available filters them. Persistence is intentionally temporary.",
        "Change exercise: if repeated withdrawal should succeed idempotently, change B's transition contract and T8's expected outcome; booking rules need not change.",
        "AI assistance is disclosed in README. Each member must explain this tradeoff and trace the code independently.",
    ]),
    ("Every rule has inspectable evidence", "Rule → responsible code → test → final result", [
        "BR1 → Domain/ValueObject/seatnumber.py : SeatNumber → T1 → PASS",
        "BR2 → Domain/Aggregate/TripAggregate/trip.py : book_seat → T2 → PASS",
        "BR3 → Trip.book_seat + booked_seats read-only view → T3 → PASS",
        "BR4 → Domain/Services/BusSchedulingPolicy.py : ensure_available → T4 → PASS; red/green logs retained",
        "BR5 → Trip + TripFullyBooked + handler + TripAdvertisement.withdraw → T5 / T7 / T8 → PASS",
        "BR6 → BookSeatService / CreateTripService + ITripRepository → T6 → PASS",
        "Evidence: evidence/test-results.txt; run python -B -m unittest discover -s test -v from the root.",
        "Conclusion: six rules, two roots, an event-connected workflow, eight tests, and 15 slides. Four confirmed members cover five roles; each member must defend the work individually.",
    ]),
]


def build_html(out):
    sections = []
    for i, (title, subtitle, lines) in enumerate(SLIDES, 1):
        body = '<ul>' + ''.join(f'<li>{escape(line)}</li>' for line in lines) + '</ul>'
        if i == 8:
            body = '''<svg viewBox="0 0 1000 270" role="img" aria-label="Interface and Infrastructure depend on Application and Domain; Application depends on Domain">
<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8" fill="#087f8c"/></marker></defs>
<g fill="#e8f3f3" stroke="#087f8c" stroke-width="2"><rect x="20" y="20" width="240" height="60" rx="8"/><rect x="380" y="20" width="240" height="60" rx="8"/><rect x="740" y="20" width="240" height="60" rx="8"/><rect x="380" y="195" width="240" height="60" rx="8"/></g>
<g fill="#182b3b" font-size="23" font-family="Segoe UI,Arial" text-anchor="middle"><text x="140" y="58">Interface</text><text x="500" y="58">Application</text><text x="860" y="58">Infrastructure</text><text x="500" y="234">Domain</text></g>
<g stroke="#087f8c" stroke-width="3" fill="none" marker-end="url(#arrow)"><path d="M260,50 H375"/><path d="M740,50 H625"/><path d="M500,80 V190"/><path d="M140,80 L375,222"/><path d="M860,80 L625,222"/></g></svg>
<ul><li>main.py composes all four areas and injects the repository implementations.</li><li>Arrows show actual static imports. Domain imports only Domain modules and the standard library.</li><li>No Layer Supertype: no common entity behaviour needs one. Repository ABCs are contracts.</li></ul>'''
        sections.append(f'<section id="slide-{i}"><header>BUS SEAT BOOKING / {i:02d}</header>'
                        f'<h1>{escape(title)}</h1><h2>{escape(subtitle)}</h2>' + body
                        + f'<footer>MAKERERE UNIVERSITY · ADVANCED PROGRAMMING <span>{i} / 15</span></footer></section>')
    out.write_text('''<!doctype html><html lang="en"><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Bus seat booking — coursework slides</title><style>
*{box-sizing:border-box}body{margin:0;background:#dce3e7;color:#182b3b;font-family:Segoe UI,Arial,sans-serif}
section{position:relative;background:#fffdf7;width:min(1120px,96vw);min-height:630px;margin:24px auto;padding:38px 54px 70px;box-shadow:0 8px 30px #1232;border-top:10px solid #087f8c;scroll-margin-top:12px}
header{font-size:12px;letter-spacing:2px;color:#087f8c;font-weight:700}h1{font-size:38px;margin:18px 0 6px;line-height:1.1}h2{font-size:18px;font-weight:500;color:#587080;margin:0 0 24px}ul{padding-left:22px;margin:0}li{font-size:18px;line-height:1.4;margin:0 0 13px;padding-left:5px}footer{position:absolute;bottom:25px;left:54px;right:54px;border-top:1px solid #cad6dc;padding-top:12px;font-size:10px;letter-spacing:1px;color:#587080}footer span{float:right}
nav{position:fixed;right:18px;bottom:12px;z-index:2;display:flex;gap:8px}button{background:#182b3b;color:white;border:0;border-radius:5px;padding:10px 14px;cursor:pointer}
@page{size:landscape;margin:0}@media print{body{background:white}section{width:100vw;height:100vh;min-height:0;margin:0;box-shadow:none;break-after:page;padding:5vh 5vw 9vh}h1{font-size:30px}li{font-size:15px;margin-bottom:10px}nav{display:none}section:last-of-type{break-after:auto}}
</style><nav aria-label="Slide navigation"><button onclick="move(-1)">Previous</button><button onclick="move(1)">Next</button><button onclick="print()">Print / PDF</button></nav>'''
                   + ''.join(sections) + '''<script>
let current=0;const slides=[...document.querySelectorAll('section')];
function move(delta){current=Math.max(0,Math.min(14,current+delta));slides[current].scrollIntoView({behavior:'smooth'});}
document.addEventListener('keydown',e=>{if(['ArrowRight','PageDown','ArrowLeft','PageUp','Home','End'].includes(e.key)){e.preventDefault();if(e.key==='Home')current=0;else if(e.key==='End')current=14;else current=Math.max(0,Math.min(14,current+(['ArrowRight','PageDown'].includes(e.key)?1:-1)));slides[current].scrollIntoView({behavior:'smooth'});}});
const observer=new IntersectionObserver(entries=>{for(const e of entries)if(e.isIntersecting)current=slides.indexOf(e.target);},{threshold:0.6});slides.forEach(s=>observer.observe(s));
</script></html>''', encoding='utf-8')


def build_pptx(out):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    def text(slide, x, y, w, h, lines, size, color):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = 'Aptos'
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(10)
        return shape

    for i, (title, subtitle, lines) in enumerate(SLIDES, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFDF7')
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, Inches(.1))
        band.fill.solid()
        band.fill.fore_color.rgb = RGBColor.from_string('087F8C')
        band.line.fill.background()
        text(slide, .65, .3, 12, .3, [f'BUS SEAT BOOKING / {i:02d}'], 10, '087F8C')
        text(slide, .65, .8, 12, .65, [title], 30, '182B3B')
        text(slide, .65, 1.55, 12, .45, [subtitle], 15, '587080')
        if i != 8:
            text(slide, .65, 2.15, 12, 4.7, ['\u2022 ' + line for line in lines], 16 if i in (3, 11, 15) else 18, '182B3B')
        else:
            for x, y, label in [(.8, 2.4, 'Interface'), (5.2, 2.4, 'Application'), (9.6, 2.4, 'Infrastructure'), (5.2, 4.5, 'Domain')]:
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.8), Inches(.75))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string('E8F3F3')
                shape.line.color.rgb = RGBColor.from_string('087F8C')
                text(slide, x+.15, y+.16, 2.5, .5, [label], 21, '182B3B')
            for x1, y1, x2, y2 in [(3.6,2.77,5.1,2.77), (9.6,2.77,8.1,2.77), (6.6,3.2,6.6,4.4), (2.2,3.2,5.1,4.87), (11,3.2,8.1,4.87)]:
                connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
                connector.line.color.rgb = RGBColor.from_string('087F8C')
                connector.line.width = Pt(2)
                arrow = OxmlElement('a:tailEnd')
                arrow.set('type', 'triangle')
                connector.line._get_or_add_ln().append(arrow)
            text(slide, .65, 5.65, 12, 1, ['main.py composes all four areas; Domain imports only Domain and the standard library.', 'No Layer Supertype: no shared entity behaviour needs one. Repository ABCs are contracts.'], 15, '587080')
        text(slide, .65, 7, 12, .3, [f'MAKERERE UNIVERSITY \u00b7 ADVANCED PROGRAMMING                                                 {i} / 15'], 9, '587080')
    deck.save(out)


if __name__ == '__main__':
    output = ROOT / 'presentation'
    output.mkdir(exist_ok=True)
    build_html(output / 'coursework-slides.html')
    build_pptx(output / 'coursework-slides.pptx')
    print('Built 15 slides in HTML and editable PowerPoint formats.')
